#!/usr/bin/env python3
"""Pre-delivery lint — deterministic deliverable check (zero cost).
Usage: pre_delivery_lint.py <docx|md|txt> [--cutoff YYYY-MM-DD]
Checks: dates past cutoff · mixed languages · missing 'did not verify'
section · risk claims without hedging · numbers without adjacent source.
The linter flags; the human decides.
"""
import re
import sys
import datetime
import pathlib


def text_of(path):
    p = pathlib.Path(path)
    if p.suffix.lower() == ".docx":
        import docx
        d = docx.Document(str(p))
        parts = [par.text for par in d.paragraphs]
        for tb in d.tables:
            for r in tb.rows:
                for c in r.cells:
                    parts.append(c.text)
        return "\n".join(parts)
    return p.read_text(encoding="utf-8", errors="ignore")


def main():
    args = sys.argv[1:]
    if not args:
        print(__doc__)
        sys.exit(1)
    path = args[0]
    cutoff = None
    if "--cutoff" in args:
        cutoff = args[args.index("--cutoff") + 1]
    txt = text_of(path)
    warnings = []

    # 1) dates past the evidence cutoff
    if cutoff:
        cut = datetime.date.fromisoformat(cutoff)
        for d, y, m, dd in re.findall(
            r"(\b(\d{1,2})[/\-\. ](\d{1,2})[/\-\. ](\d{2,4})\b)", txt
        ):
            try:
                dd4 = int(dd) if int(dd) > 31 else 2000 + int(dd)
                dt = datetime.date(dd4, int(m), int(y)) if int(m) <= 12 else None
                if dt and dt > cut:
                    warnings.append(f"DATE>CUTOFF: {d} > {cut}")
            except Exception:
                pass
        for iso in re.findall(r"\b(20\d{2})-(\d{2})-(\d{2})\b", txt):
            try:
                dt = datetime.date(int(iso[0]), int(iso[1]), int(iso[2]))
                if dt > cut:
                    warnings.append(f"DATE>CUTOFF: {'-'.join(iso)}")
            except Exception:
                pass

    # 2) mixed language (heuristic: typical FR markers in a mostly-EN doc)
    fr_markers = len(
        re.findall(
            r"\b(donc|toutefois|neanmoins|afin|ainsi|mise a jour|deplaces|ecoles)\b",
            txt,
            re.I,
        )
    )
    en_markers = len(
        re.findall(
            r"\b(the|and|with|from|reported|according|displaced|schools)\b", txt, re.I
        )
    )
    if en_markers > 20 and fr_markers > en_markers * 0.25:
        warnings.append(f"MIXED LANGUAGE: FR markers={fr_markers} vs EN={en_markers}")

    # 3) 'not verified' section present
    if not re.search(r"(not verified|did NOT verify|unverified|to be verified)", txt, re.I):
        warnings.append("MISSING SECTION: 'did not verify' / limits not declared")

    # 4) risk claims without nearby hedging
    for kw in (
        r"\bfamine\b",
        r"\bmassacre\b",
        r"\bpredict\w*\b",
        r"\bwill (collapse|fall)\b",
    ):
        for m in re.finditer(kw, txt, re.I):
            ctx = txt[max(0, m.start() - 120): m.end() + 120]
            if not re.search(
                r"(risk|claim|alleged|report|estimat|according|phase|unconfirm|uncertain|plausible|scenario|not predictions)",
                ctx,
                re.I,
            ):
                snippet = ctx.strip().replace("\n", " ")[:60]
                warnings.append(f"RISK CLAIM without hedging: ...{snippet}...")
                break  # one per pattern

    # 5) bare numbers without adjacent source (heuristic)
    bare = 0
    for m in re.finditer(r"\b\d{2,3}[ ,]?\d{3}\b", txt):
        ctx = txt[max(0, m.start() - 100): m.end() + 100]
        if not re.search(
            r"(source|per|reported|estimate|according|UN|NGO|outlet|survey)",
            ctx,
            re.I,
        ):
            bare += 1
    if bare:
        warnings.append(
            f"{bare} number(s) without adjacent source/reference (check manually)"
        )

    if warnings:
        print(f"LINT {path} — {len(warnings)} warning(s):")
        for w in warnings:
            print("  ⚠", w)
        sys.exit(1)
    else:
        print(f"LINT {path} — OK (no warnings)")


if __name__ == "__main__":
    main()
